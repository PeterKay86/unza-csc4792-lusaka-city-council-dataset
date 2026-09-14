import csv
import hashlib
import re
from pathlib import Path

import openpyxl
import pdfplumber


PROJECT_ROOT = Path(__file__).resolve().parents[2]

SOURCE_DIRECTORY = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "cdf"
    / "source_documents"
)

TEXT_DIRECTORY = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "cdf"
    / "extracted_text"
)

TABLE_DIRECTORY = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "cdf"
    / "extracted_tables"
)

MANIFEST_FILE = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "cdf"
    / "db-unza26-csc4792-lusaka_city_council_cdf_source_manifest.csv"
)


def safe_name(value):
    """Create a safe lowercase name for an output file."""

    value = value.lower().strip()
    value = re.sub(r"[^a-z0-9]+", "_", value)

    return value.strip("_")


def calculate_sha256(file_path):
    """Calculate a hash used to identify duplicate files."""

    sha256 = hashlib.sha256()

    with file_path.open("rb") as source_file:
        for block in iter(lambda: source_file.read(65536), b""):
            sha256.update(block)

    return sha256.hexdigest()


def extract_pdf(file_path):
    """Extract text from every page of a PDF."""

    output_path = TEXT_DIRECTORY / f"{safe_name(file_path.stem)}.txt"

    page_texts = []

    with pdfplumber.open(file_path) as pdf:

        page_count = len(pdf.pages)

        for page_number, page in enumerate(pdf.pages, start=1):

            page_text = page.extract_text() or ""

            page_texts.append(
                f"\n===== PAGE {page_number} =====\n{page_text}"
            )

    complete_text = "\n".join(page_texts).strip()

    output_path.write_text(
        complete_text,
        encoding="utf-8",
    )

    extraction_status = (
        "Extracted"
        if complete_text
        else "No extractable text; OCR may be required"
    )

    return {
        "page_count": page_count,
        "sheet_count": "",
        "extracted_record_count": "",
        "output_file": str(output_path.relative_to(PROJECT_ROOT)),
        "extraction_status": extraction_status,
    }


def extract_excel(file_path):
    """Extract every Excel worksheet as a pipe-separated CSV."""

    workbook = openpyxl.load_workbook(
        file_path,
        data_only=True,
        read_only=True,
    )

    workbook_folder = (
        TABLE_DIRECTORY / safe_name(file_path.stem)
    )

    workbook_folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    total_rows = 0
    output_files = []

    for worksheet in workbook.worksheets:

        output_path = (
            workbook_folder
            / f"{safe_name(worksheet.title)}.csv"
        )

        row_count = 0

        with output_path.open(
            "w",
            encoding="utf-8-sig",
            newline="",
        ) as output_file:

            writer = csv.writer(
                output_file,
                delimiter="|",
                quoting=csv.QUOTE_MINIMAL,
            )

            for row in worksheet.iter_rows(values_only=True):

                cleaned_row = [
                    "" if value is None else value
                    for value in row
                ]

                if any(value != "" for value in cleaned_row):
                    writer.writerow(cleaned_row)
                    row_count += 1

        total_rows += row_count

        output_files.append(
            str(output_path.relative_to(PROJECT_ROOT))
        )

    return {
        "page_count": "",
        "sheet_count": len(workbook.sheetnames),
        "extracted_record_count": total_rows,
        "output_file": "; ".join(output_files),
        "extraction_status": "Extracted",
    }


def create_manifest():
    """Extract all supported source documents and create a manifest."""

    TEXT_DIRECTORY.mkdir(
        parents=True,
        exist_ok=True,
    )

    TABLE_DIRECTORY.mkdir(
        parents=True,
        exist_ok=True,
    )

    source_files = sorted(
        file_path
        for file_path in SOURCE_DIRECTORY.iterdir()
        if file_path.suffix.lower() in {".pdf", ".xlsx"}
    )

    if not source_files:
        raise FileNotFoundError(
            f"No PDF or Excel files found in {SOURCE_DIRECTORY}"
        )

    manifest_rows = []
    known_hashes = {}

    for file_number, file_path in enumerate(source_files, start=1):

        print(
            f"Processing {file_number} of "
            f"{len(source_files)}: {file_path.name}"
        )

        file_hash = calculate_sha256(file_path)

        duplicate_of = known_hashes.get(file_hash, "")

        if not duplicate_of:
            known_hashes[file_hash] = file_path.name

        try:
            if file_path.suffix.lower() == ".pdf":
                result = extract_pdf(file_path)
                document_type = "PDF"

            else:
                result = extract_excel(file_path)
                document_type = "Excel workbook"

            if duplicate_of:
                result["extraction_status"] = (
                    f"Duplicate of {duplicate_of}"
                )

        except Exception as error:
            result = {
                "page_count": "",
                "sheet_count": "",
                "extracted_record_count": "",
                "output_file": "",
                "extraction_status": f"Failed: {error}",
            }

            document_type = file_path.suffix.lower().replace(".", "").upper()

        manifest_rows.append(
            {
                "source_id": f"LCC-SOURCE-{file_number:03d}",
                "file_name": file_path.name,
                "document_type": document_type,
                "file_size_bytes": file_path.stat().st_size,
                "sha256": file_hash,
                "duplicate_of": duplicate_of,
                "page_count": result["page_count"],
                "sheet_count": result["sheet_count"],
                "extracted_record_count": result[
                    "extracted_record_count"
                ],
                "output_file": result["output_file"],
                "extraction_status": result["extraction_status"],
            }
        )

    fieldnames = [
        "source_id",
        "file_name",
        "document_type",
        "file_size_bytes",
        "sha256",
        "duplicate_of",
        "page_count",
        "sheet_count",
        "extracted_record_count",
        "output_file",
        "extraction_status",
    ]

    with MANIFEST_FILE.open(
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as manifest_file:

        writer = csv.DictWriter(
            manifest_file,
            fieldnames=fieldnames,
            delimiter="|",
        )

        writer.writeheader()
        writer.writerows(manifest_rows)

    print()
    print(f"Documents processed: {len(source_files)}")
    print(f"Manifest saved to: {MANIFEST_FILE}")
    print(f"PDF text saved to: {TEXT_DIRECTORY}")
    print(f"Excel tables saved to: {TABLE_DIRECTORY}")


if __name__ == "__main__":
    create_manifest()